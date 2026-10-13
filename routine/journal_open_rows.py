"""List the journal's Open option rows and tab names as JSON, for the daily IBKR routine.

usage: python3 -I routine/journal_open_rows.py journal.xlsx > open_rows.json

Wheel tabs are the ones whose row 7 reads Date Opened / Type / ... / Status;
their option block runs from row 8 to two rows above the "... SHARES" title.
"""
import json
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")


def iso(v):
    return v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else v


def wheel_last_row(ws):
    for r in range(8, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and "SHARES" in a.upper():
            return r - 2
    return ws.max_row


# BPS tracker columns are found by their row-5 header (same patterns as the Apps Script)
BPS_HDR = {"date": r"^Trade Date", "status": r"^Status$", "symbol": r"^Stock",
           "expiry": r"^Expiration", "strategy": r"^Strategy$", "qty": r"^No\.?$",
           "short_strike": r"^Strike ?1", "long_strike": r"^Strike ?2", "entry": r"^Entry Price"}


def bps_cols(ws):
    cols = {}
    for c in range(1, ws.max_column + 1):
        v = str(ws.cell(5, c).value or "").strip()
        for k, pat in BPS_HDR.items():
            if k not in cols and v and re.search(pat, v, re.I):
                cols[k] = c
    return cols


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {"wheel_tabs": [], "open_wheel_rows": [], "open_bps_rows": [], "bps_tab": None}
    for ws in wb:
        h = [ws.cell(7, c).value for c in (1, 2, 15)]
        if h == ["Date Opened", "Type", "Status"]:
            out["wheel_tabs"].append(ws.title)
            for r in range(8, wheel_last_row(ws) + 1):
                if ws.cell(r, 15).value != "Open":
                    continue
                out["open_wheel_rows"].append({
                    "tab": ws.title, "row": r, "date_opened": iso(ws.cell(r, 1).value),
                    "type": ws.cell(r, 2).value, "symbol": ws.cell(r, 3).value,
                    "expiry": iso(ws.cell(r, 4).value), "strike": ws.cell(r, 5).value,
                    "qty": ws.cell(r, 7).value, "open_price": ws.cell(r, 8).value,
                })
        elif ws.title == "Bull PUT SPREAD TRACKER":
            out["bps_tab"] = ws.title
            cols = bps_cols(ws)
            for r in range(6, ws.max_row + 1):
                if ws.cell(r, cols["status"]).value != "Open":
                    continue
                row = {k: ws.cell(r, c).value for k, c in cols.items()}
                row.update(tab=ws.title, row=r, date=iso(row.get("date")),
                           expiry=iso(row.get("expiry")))
                out["open_bps_rows"].append(row)
    json.dump(out, sys.stdout, indent=1, default=str)


if __name__ == "__main__":
    main(sys.argv[1])
